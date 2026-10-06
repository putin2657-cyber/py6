#!/usr/bin/env python3
"""Конспект markdown → настоящий .docx для Word / LibreOffice Writer."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = Path("/workspace/ucheba/PLAN_3KR_4DNYA.md")
OUT = Path("/workspace/ucheba/Konspekt_3KR_chislennye.docx")


def set_run_font(run, name="Times New Roman", size=12, bold=False, italic=False, mono=False):
    if mono:
        name = "Courier New"
        size = 10.5
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    rPr = run._element.get_or_add_rPr()
    lang = rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rPr.append(lang)
    lang.set(qn("w:val"), "ru-RU")
    lang.set(qn("w:eastAsia"), "ru-RU")
    if rPr.find(qn("w:noProof")) is None:
        rPr.append(OxmlElement("w:noProof"))


def add_inline(p, text: str, base_size=12, mono_size=10.5):
    # split **bold** and `code`
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            set_run_font(r, size=base_size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            set_run_font(r, size=mono_size, mono=True)
        else:
            r = p.add_run(part)
            set_run_font(r, size=base_size)


def is_table_sep(line: str) -> bool:
    s = line.strip()
    return bool(re.match(r"^\|?[\s:-]+\|[\s|:-]+$", s))


def parse_row(line: str) -> list[str]:
    parts, buf, in_code = [], [], False
    for ch in line.strip():
        if ch == "`":
            in_code = not in_code
            buf.append(ch)
        elif ch == "|" and not in_code:
            parts.append("".join(buf))
            buf = []
        else:
            buf.append(ch)
    parts.append("".join(buf))
    cells = [c.strip() for c in parts]
    while cells and cells[0] == "":
        cells.pop(0)
    while cells and cells[-1] == "":
        cells.pop()
    return cells


def add_table(doc, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=cols)
    tbl.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = tbl.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            val = row[j] if j < len(row) else ""
            add_inline(p, val, base_size=11)
            if i == 0:
                for r in p.runs:
                    r.bold = True


def convert():
    text = SRC.read_text(encoding="utf-8").replace("- [ ]", "☐")
    lines = text.splitlines()

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(1.5)
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.5)

    styles = doc.styles["Normal"]
    styles.font.name = "Times New Roman"
    styles.font.size = Pt(12)
    styles._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = styles.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        raw = line.rstrip()
        s = raw.strip()

        if not s:
            i += 1
            continue

        if s == "---":
            i += 1
            continue

        if s.startswith("# "):
            p = doc.add_heading(s[2:].strip(), level=1)
            for r in p.runs:
                set_run_font(r, size=18, bold=True)
            i += 1
            continue
        if s.startswith("## "):
            p = doc.add_heading(s[3:].strip(), level=2)
            for r in p.runs:
                set_run_font(r, size=14, bold=True)
                r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
            i += 1
            continue
        if s.startswith("### "):
            p = doc.add_heading(s[4:].strip(), level=3)
            for r in p.runs:
                set_run_font(r, size=13, bold=True)
            i += 1
            continue

        if s.startswith("```"):
            buf = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1  # closing ```
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run("\n".join(buf))
            set_run_font(r, mono=True)
            continue

        if s.startswith("|") and i + 1 < n and is_table_sep(lines[i + 1]):
            rows = [parse_row(s)]
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append(parse_row(lines[i]))
                i += 1
            add_table(doc, rows)
            doc.add_paragraph("")
            continue

        if s.startswith("> "):
            buf = [s[2:]]
            i += 1
            while i < n and lines[i].strip().startswith(">"):
                t = lines[i].strip()
                buf.append(t[2:] if t.startswith("> ") else t[1:].lstrip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.space_before = Pt(4)
            add_inline(p, " ".join(buf))
            for r in p.runs:
                r.italic = True
            continue

        m_ul = re.match(r"^[-*] (.+)$", s)
        if m_ul:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, m_ul.group(1))
            i += 1
            continue

        m_ol = re.match(r"^(\d+)\. (.+)$", s)
        if m_ol:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, m_ol.group(2))
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline(p, s)
        i += 1

    core = doc.core_properties
    core.title = "Конспект: 3 КР по численным методам"
    core.language = "ru-RU"
    sett = doc.settings.element
    for tag in ("w:hideSpellingErrors", "w:hideGrammaticalErrors"):
        if sett.find(qn(tag)) is None:
            sett.append(OxmlElement(tag))
    doc.save(OUT)
    print(OUT, OUT.stat().st_size)


if __name__ == "__main__":
    convert()
